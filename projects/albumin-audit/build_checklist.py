"""ساخت فایل Excel چک‌لیست ممیزی پرونده‌های آلبومین (بر اساس فرم درخواست آلبومین بیمارستان)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent / "albumin_audit_checklist.xlsx"
N_ROWS = 200
FIRST = 3
LAST = FIRST + N_ROWS - 1

F = "Arial"
TITLE = Font(name=F, size=14, bold=True)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E7EEF7")
HFILL = PatternFill("solid", fgColor="1F4E78")
GROUP = PatternFill("solid", fgColor="D9D9D9")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="right", readingOrder=2)
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center", readingOrder=2)

wb = Workbook()

# ------------------------------------------------------------------ اندیکاسیون‌ها
ref = wb.active
ref.title = "اندیکاسیون‌ها"
ref.sheet_view.rightToLeft = True
ref["A1"] = "جدول مرجع اندیکاسیون‌ها (برگرفته از فرم درخواست آلبومین بیمارستان)"
ref["A1"].font = TITLE
ref_head = ["کد", "طبقه", "اندیکاسیون", "معیار الزامی که باید در پرونده مستند باشد", "دوز طبق فرم",
            "سقف ویال ۲۰٪ کل دوره (برای ممیزی)", "نقطه پایان"]
INDIC = [
    ("A1", "الف", "پاراسنتز حجیم در سیروز", "خروج ≥ ۵ لیتر آسیت (یا هر حجم با کراتینین > ۱٫۵)",
     "۶ تا ۸ گرم به ازای هر لیتر", 8, "پس از اتمام پاراسنتز"),
    ("A2", "الف", "پریتونیت باکتریایی خودبه‌خودی (SBP)",
     "PMN مایع آسیت ≥ ۲۵۰ و یکی از: کراتینین > ۱، BUN > ۳۰، بیلی‌روبین > ۴",
     "روز ۱: ۱٫۵ g/kg، روز ۳: ۱ g/kg (حداکثر ۱۰۰ گرم در روز)", 20, "فقط روز ۱ و ۳"),
    ("A3", "الف", "سندرم هپاتورنال (HRS-AKI)",
     "سیروز + کراتینین > ۱٫۵، بدون شوک، عفونت فعال یا داروی نفروتوکسیک",
     "روز ۱ و ۲: ۱ g/kg، سپس ۲۰ تا ۴۰ گرم روزانه با تنگ‌کننده عروقی", 68,
     "بهبود کراتینین یا حداکثر ۱۴ روز"),
    ("A4", "الف", "پلاسمافرز", "دستور پزشک متخصص و پروتکل آفرزیس",
     "آلبومین ۵٪ برای جایگزینی ۷۰ تا ۸۰٪ حجم پلاسمای خارج‌شده", None, "پایان هر جلسه"),
    ("B1", "ب", "احیای حجم پس از جراحی قلب",
     "≥ ۳ لیتر کریستالوئید بولوس در ۲۴ ساعت بدون پاسخ همودینامیک (مایع حین عمل حساب نمی‌شود)",
     "آلبومین ۵٪ یا معادل آن", 2, "پاسخ همودینامیک یا ۲۴ ساعت"),
    ("B2", "ب", "رزکسیون بیش از ۴۰٪ کبد",
     "آلبومین سرم < ۲٫۵ و حجم داخل عروقی ناکافی با کریستالوئید", "۲۰ تا ۲۵ گرم روزانه", 12,
     "آلبومین ≥ ۲٫۵ یا حداکثر ۴ روز"),
    ("B3", "ب", "شوک سپتیک مقاوم",
     "نیاز به وازوپرسور پس از ≥ ۳۰ mL/kg کریستالوئید؛ امضای متخصص ICU", "حداکثر ۲۰ گرم در روز", 4,
     "حداکثر ۲ روز"),
    ("C1", "ج", "سندرم نفروتیک شدید با آناسارکا یا ادم ریه",
     "پروتئینوری > ۳ g/day و مقاومت به فوروزماید (≥ ۱۶۰ mg بولوس + ≥ ۲۰ mg/h به مدت ۸ ساعت)",
     "۲۰ تا ۲۵ گرم همراه با دیورتیک", 3, "دیورز کافی؛ تکرار با تأیید نفرولوژی"),
    ("C2", "ج", "آسیت یا ادم مقاوم به دیورتیک در سیروز",
     "حداکثر دوز دیورتیک (اسپیرونولاکتون ۴۰۰ + فوروزماید ۱۶۰) + آلبومین < ۲ g/dL", "۲۰ گرم", 2,
     "فقط روز ۱"),
    ("N1", "بدون اندیکاسیون", "ترومای مغزی (TBI)", "منع مصرف", "", 0, ""),
    ("N2", "بدون اندیکاسیون", "اصلاح آلبومین پایین بدون علامت", "", "", 0, ""),
    ("N3", "بدون اندیکاسیون", "احیای حجم روتین در ICU", "", "", 0, ""),
    ("N4", "بدون اندیکاسیون", "سپسیس بدون شوک", "", "", 0, ""),
    ("N5", "بدون اندیکاسیون", "سوءتغذیه، سوءجذب، انتروپاتی دفع پروتئین", "", "", 0, ""),
    ("N6", "بدون اندیکاسیون", "افت فشار حین دیالیز / بهبود اولترافیلتراسیون", "", "", 0, ""),
    ("N7", "بدون اندیکاسیون", "ترومای عمده، شوک هموراژیک، ARDS", "", "", 0, ""),
    ("N8", "بدون اندیکاسیون", "انسفالوپاتی کبدی، پانکراتیت، نارسایی قلبی، آنمی", "", "", 0, ""),
    ("N9", "بدون اندیکاسیون", "هیپوآلبومینمی پیش از جراحی", "", "", 0, ""),
    ("X1", "خارج از جدول", "سایر موارد خارج از جدول",
     "مشاوره فارماکوتراپی و تأیید رئیس کمیته دارو و درمان یا معاون درمان", "", None, ""),
    ("U0", "نامشخص", "اندیکاسیون از پرونده قابل تشخیص نیست", "", "", 0, ""),
]
for j, h in enumerate(ref_head, 1):
    c = ref.cell(row=2, column=j, value=h)
    c.font, c.fill, c.alignment, c.border = HEAD, HFILL, CENTER, BOX
for i, row in enumerate(INDIC, 3):
    for j, v in enumerate(row, 1):
        c = ref.cell(row=i, column=j, value=v)
        c.font, c.alignment, c.border = BODY, WRAP, BOX
REF_LAST = 2 + len(INDIC)
for col, w in zip("ABCDEFG", [7, 14, 34, 52, 36, 16, 26]):
    ref.column_dimensions[col].width = w
ref["F4"].comment = Comment("مشتق از فرم: روز ۱ و ۳، هر روز حداکثر ۱۰ ویال.", "ناظر دارویی")
ref["F5"].comment = Comment("مشتق از فرم: روز ۱ و ۲ هر کدام ۱۰ ویال + ۱۲ روز × ۴ ویال = ۶۸.", "ناظر دارویی")
ref["F6"].comment = Comment("به تعداد جلسه و حجم بستگی دارد؛ ممیز سقف را دستی وارد می‌کند.", "ناظر دارویی")
ref["F8"].comment = Comment("مشتق از فرم: ۳ ویال در روز × حداکثر ۴ روز.", "ناظر دارویی")
ref["F9"].comment = Comment("مشتق از فرم: ۲ ویال در روز × حداکثر ۲ روز.", "ناظر دارویی")
note = ref.cell(row=REF_LAST + 2, column=1,
                value="سقف کل دوره از ستون‌های دوز و نقطه پایان همین فرم مشتق شده و فقط برای ممیزی است. "
                      "اگر فرم در کمیته اصلاح شد، این جدول هم به‌روز شود.")
note.font = Font(name=F, size=9, italic=True)

# ------------------------------------------------------------------ ممیزی
ws = wb.create_sheet("ممیزی")
ws.sheet_view.rightToLeft = True
ws["A1"] = "ثبت ممیزی پرونده‌های آلبومین — هر ردیف یک بستری"
ws["A1"].font = TITLE

YN = "بله/خیر/نامرتبط"
# (عنوان، نوع: in/calc، عرض، اعتبارسنجی یا فرمول)
COLS = [
    ("ردیف", "calc", 6, "row"),
    ("شماره پرونده", "in", 14, None),
    ("بخش", "in", 12, "ward"),
    ("بیمه پایه", "in", 13, "ins"),
    ("تاریخ پذیرش", "in", 12, None),
    ("وزن (kg)", "in", 8, None),
    ("کد اندیکاسیون (از پرونده)", "in", 11, "code"),
    ("اندیکاسیون", "calc", 26, "indic"),
    ("طبقه", "calc", 12, "cat"),
    ("فرم درخواست کامل و امضاشده", "in", 11, YN),
    ("اندیکاسیون در دستور پزشک", "in", 11, YN),
    ("معیار الزامی در پرونده مستند است", "in", 12, YN),
    ("آزمایش‌های الزامی فرم ثبت شده", "in", 11, YN),
    ("دوز به گرم و نقطه پایان ذکر شده", "in", 11, YN),
    ("تأیید داروساز", "in", 10, YN),
    ("تأیید موارد خارج از معیار", "in", 11, YN),
    ("ثبت تزریق همه ویال‌ها در گزارش پرستاری", "in", 12, YN),
    ("ویال تحویلی (HIS)", "in", 9, None),
    ("ویال تزریق‌شده (پرستاری)", "in", 9, None),
    ("ویال برگشتی", "in", 9, None),
    ("سقف مرجع جدول", "calc", 9, "cap"),
    ("سقف برای این بیمار (اختیاری)", "in", 10, None),
    ("قیمت خرید هر ویال (ریال)", "in", 13, None),
    ("ویال مازاد بر سقف", "calc", 9, "excess"),
    ("ویال بدون سرنوشت (تحویلی − تزریق − برگشت)", "calc", 12, "unacc"),
    ("وضعیت اندیکاسیون", "calc", 18, "status"),
    ("امتیاز مستندات", "calc", 9, "score"),
    ("ارزش مصرف نامطابق (ریال)", "calc", 15, "value"),
    ("سهم بیمه ثبت‌شده (ریال)", "in", 14, None),
    ("سهم بیمار ثبت‌شده (ریال)", "in", 14, None),
    ("پیام بیمه", "calc", 20, "insflag"),
    ("یادداشت ممیز / اقدام", "in", 28, None),
]
L = {name: get_column_letter(i) for i, (name, *_) in enumerate(COLS, 1)}
k = {key: L[name] for name, _, _, key in COLS if key and key != YN}
k["file"] = L["شماره پرونده"]
k["ins_col"] = L["بیمه پایه"]
k["code_col"] = L["کد اندیکاسیون (از پرونده)"]
k["crit"] = L["معیار الزامی در پرونده مستند است"]
k["outapp"] = L["تأیید موارد خارج از معیار"]
k["deliv"] = L["ویال تحویلی (HIS)"]
k["given"] = L["ویال تزریق‌شده (پرستاری)"]
k["ret"] = L["ویال برگشتی"]
k["cap_in"] = L["سقف برای این بیمار (اختیاری)"]
k["price"] = L["قیمت خرید هر ویال (ریال)"]
k["ins_amt"] = L["سهم بیمه ثبت‌شده (ریال)"]
yn_first = L["فرم درخواست کامل و امضاشده"]
yn_last = L["ثبت تزریق همه ویال‌ها در گزارش پرستاری"]
REFR = f"'اندیکاسیون‌ها'!$A$3:$A${REF_LAST}"


def formula(key, r):
    f = k["file"]
    if key == "row":
        return f'=IF({f}{r}="","",ROW()-{FIRST - 1})'
    if key == "indic":
        return (f'=IF({k["code_col"]}{r}="","",IFERROR(INDEX(\'اندیکاسیون‌ها\'!$C$3:$C${REF_LAST},'
                f'MATCH({k["code_col"]}{r},{REFR},0)),"کد نامعتبر"))')
    if key == "cat":
        return (f'=IF({k["code_col"]}{r}="","",IFERROR(INDEX(\'اندیکاسیون‌ها\'!$B$3:$B${REF_LAST},'
                f'MATCH({k["code_col"]}{r},{REFR},0)),""))')
    if key == "cap":
        return (f'=IF({k["code_col"]}{r}="","",IFERROR(INDEX(\'اندیکاسیون‌ها\'!$F$3:$F${REF_LAST},'
                f'MATCH({k["code_col"]}{r},{REFR},0)),""))')
    if key == "excess":
        cap = f'IF({k["cap_in"]}{r}<>"",{k["cap_in"]}{r},{k["cap"]}{r})'
        return (f'=IF(OR({f}{r}="",{cap}=""),"",'
                f'MAX(0,N({k["deliv"]}{r})-N({k["ret"]}{r})-{cap}))')
    if key == "unacc":
        return f'=IF({f}{r}="","",N({k["deliv"]}{r})-N({k["given"]}{r})-N({k["ret"]}{r}))'
    if key == "status":
        c, cat, ex = k["code_col"], k["cat"], k["excess"]
        return (f'=IF({f}{r}="","",IF(OR({c}{r}="",{c}{r}="U0"),"اندیکاسیون نامشخص",'
                f'IF({cat}{r}="بدون اندیکاسیون","بدون اندیکاسیون",'
                f'IF(AND({cat}{r}="خارج از جدول",{k["outapp"]}{r}<>"بله"),"خارج از جدول بدون تأیید",'
                f'IF({k["crit"]}{r}<>"بله","معیار مستند نشده",'
                f'IF(N({ex}{r})>0,"مطابق، دوز مازاد","مطابق"))))))')
    if key == "score":
        rng = f"{yn_first}{r}:{yn_last}{r}"
        return (f'=IF({f}{r}="","",IFERROR(COUNTIF({rng},"بله")/'
                f'(COUNTIF({rng},"بله")+COUNTIF({rng},"خیر")),""))')
    if key == "value":
        s, p = k["status"], k["price"]
        net = f'(N({k["deliv"]}{r})-N({k["ret"]}{r}))'
        return (f'=IF({f}{r}="","",IF({s}{r}="مطابق",0,IF({s}{r}="مطابق، دوز مازاد",'
                f'N({k["excess"]}{r})*N({p}{r}),{net}*N({p}{r}))))')
    if key == "insflag":
        s, amt, ins = k["status"], k["ins_amt"], k["ins_col"]
        ok = f'OR({s}{r}="مطابق",{s}{r}="مطابق، دوز مازاد")'
        return (f'=IF({f}{r}="","",IF(AND(N({amt}{r})>0,NOT({ok})),"ریسک کسورات",'
                f'IF(AND({ok},N({amt}{r})=0,{ins}{r}<>"آزاد",{ins}{r}<>""),'
                f'"بررسی امکان دریافت از بیمه","")))')
    raise KeyError(key)


for j, (name, kind, width, spec) in enumerate(COLS, 1):
    col = get_column_letter(j)
    c = ws.cell(row=2, column=j, value=name)
    c.font, c.fill, c.alignment, c.border = HEAD, HFILL, CENTER, BOX
    ws.column_dimensions[col].width = width
    for r in range(FIRST, LAST + 1):
        cell = ws.cell(row=r, column=j)
        cell.font, cell.border, cell.alignment = BODY, BOX, CENTER
        if kind == "calc":
            cell.value = formula(spec, r)
            cell.fill = CALC
        else:
            cell.fill = INPUT
ws.row_dimensions[2].height = 75
ws.freeze_panes = f"{L['اندیکاسیون']}{FIRST}"

for key in ["price", "ins_amt"]:
    for r in range(FIRST, LAST + 1):
        ws[f"{k[key]}{r}"].number_format = "#,##0"
ins_bal = L["سهم بیمار ثبت‌شده (ریال)"]
for r in range(FIRST, LAST + 1):
    ws[f"{ins_bal}{r}"].number_format = "#,##0"
    ws[f"{k['value']}{r}"].number_format = "#,##0;-#,##0;-"
    ws[f"{k['score']}{r}"].number_format = "0%"

WARDS = ["ICU", "I.C.U OH", "C.C.U", "بخش 2", "بخش 3A", "بخش 3B", "بخش 4A", "بخش 4B", "سایر"]
INS = ["تامین اجتماعی", "سلامت", "نیروهای مسلح", "سایر بیمه‌ها", "آزاد"]
lists = wb.create_sheet("فهرست‌ها")
for i, v in enumerate(WARDS, 1):
    lists.cell(row=i, column=1, value=v)
for i, v in enumerate(INS, 1):
    lists.cell(row=i, column=2, value=v)
for i, v in enumerate(["بله", "خیر", "نامرتبط"], 1):
    lists.cell(row=i, column=3, value=v)
lists.sheet_state = "hidden"


def add_dv(src, col):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    dv.error, dv.errorTitle = "از فهرست انتخاب کنید", "مقدار نامعتبر"
    ws.add_data_validation(dv)
    dv.add(f"{col}{FIRST}:{col}{LAST}")


add_dv(f"'فهرست‌ها'!$A$1:$A${len(WARDS)}", L["بخش"])
add_dv(f"'فهرست‌ها'!$B$1:$B${len(INS)}", L["بیمه پایه"])
add_dv(REFR, k["code_col"])
for name, kind, _, spec in COLS:
    if spec == YN:
        add_dv("'فهرست‌ها'!$C$1:$C$3", L[name])
num = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
num.error = "عدد صحیح صفر یا بیشتر"
ws.add_data_validation(num)
for key in ["deliv", "given", "ret", "cap_in"]:
    num.add(f"{k[key]}{FIRST}:{k[key]}{LAST}")

ws[f"{L['معیار الزامی در پرونده مستند است']}2"].comment = Comment(
    "مبنای قضاوت «مطابق»: یافته یا آزمایشی که معیار ستون D جدول اندیکاسیون‌ها را نشان دهد، "
    "پیش از تجویز در پرونده ثبت شده باشد.", "ناظر دارویی")
ws[f"{k['unacc']}2"].comment = Comment(
    "باید صفر باشد. عدد مثبت یعنی ویالی که تحویل شده اما نه تزریقش ثبت شده و نه برگشته است.", "ناظر دارویی")
ws[f"{k['value']}2"].comment = Comment(
    "برای وضعیت‌های نامطابق: (تحویلی − برگشتی) × قیمت. برای «مطابق، دوز مازاد»: فقط ویال‌های مازاد × قیمت.",
    "ناظر دارویی")

# ------------------------------------------------------------------ خلاصه
sm = wb.create_sheet("خلاصه")
sm.sheet_view.rightToLeft = True
sm["A1"] = "خلاصه نتایج ممیزی آلبومین"
sm["A1"].font = TITLE
A = "'ممیزی'!"


def rng(col):
    return f"{A}${col}${FIRST}:${col}${LAST}"


r = 3


def put(label, f, fmt=None, bold=False):
    global r
    a = sm.cell(row=r, column=1, value=label)
    b = sm.cell(row=r, column=2, value=f)
    a.font = BOLD if bold else BODY
    b.font = BOLD if bold else BODY
    a.alignment = WRAP
    a.border = b.border = BOX
    if fmt:
        b.number_format = fmt
    r += 1


def section(title):
    global r
    r += 1
    c = sm.cell(row=r, column=1, value=title)
    c.font, c.fill = BOLD, GROUP
    sm.cell(row=r, column=2).fill = GROUP
    r += 1


section("کلیات")
put("تعداد پرونده ممیزی‌شده", f"=COUNTA({rng(k['file'])})", "0", bold=True)
total = f"B{r - 1}"
STATUSES = ["مطابق", "مطابق، دوز مازاد", "معیار مستند نشده", "بدون اندیکاسیون",
            "خارج از جدول بدون تأیید", "اندیکاسیون نامشخص"]
section("وضعیت اندیکاسیون (تعداد و درصد)")
sm.cell(row=r - 1, column=3, value="درصد").font = BOLD
st_rows = {}
for s in STATUSES:
    put(s, f'=COUNTIF({rng(k["status"])},"{s}")', "0")
    sm.cell(row=r - 1, column=3, value=f"=IFERROR(B{r - 1}/{total},0)").number_format = "0.0%"
    st_rows[s] = r - 1
put("نرخ مصرف مطابق اندیکاسیون (شامل دوز مازاد)",
    f"=IFERROR((B{st_rows['مطابق']}+B{st_rows['مطابق، دوز مازاد']})/{total},0)", "0.0%", bold=True)

section("ویال و ریال")
put("ویال تحویلی", f"=SUM({rng(k['deliv'])})", "#,##0")
put("ویال تزریق‌شده (ثبت پرستاری)", f"=SUM({rng(k['given'])})", "#,##0")
put("ویال برگشتی", f"=SUM({rng(k['ret'])})", "#,##0")
put("ویال بدون سرنوشت", f"=SUM({rng(k['unacc'])})", "#,##0", bold=True)
put("ویال مازاد بر سقف (فقط در موارد با اندیکاسیون مطابق)",
    f'=SUMIF({rng(k["status"])},"مطابق، دوز مازاد",{rng(k["excess"])})', "#,##0")
put("ارزش مصرف نامطابق (ریال)", f"=SUM({rng(k['value'])})", "#,##0", bold=True)
put("میانگین امتیاز مستندات", f"=IFERROR(AVERAGE({rng(k['score'])}),0)", "0%")

section("کامل بودن مستندات (درصد «بله» از موارد مرتبط)")
for name, kind, _, spec in COLS:
    if spec == YN:
        c = rng(L[name])
        put(name, f'=IFERROR(COUNTIF({c},"بله")/(COUNTIF({c},"بله")+COUNTIF({c},"خیر")),0)', "0%")

section("بیمه")
put("پرونده با «ریسک کسورات»", f'=COUNTIF({rng(k["insflag"])},"ریسک کسورات")', "0")
put("پرونده مطابق که سهم بیمه ندارد (بررسی امکان دریافت)",
    f'=COUNTIF({rng(k["insflag"])},"بررسی امکان دریافت از بیمه")', "0")
put("جمع سهم بیمه ثبت‌شده (ریال)", f"=SUM({rng(k['ins_amt'])})", "#,##0")
put("جمع سهم بیمار ثبت‌شده (ریال)", f"=SUM({rng(ins_bal)})", "#,##0")

section("به تفکیک بخش")
hdr = ["بخش", "پرونده", "مطابق (با دوز مازاد)", "نرخ مطابق", "ارزش نامطابق (ریال)", "ویال بدون سرنوشت"]
for j, h in enumerate(hdr, 1):
    c = sm.cell(row=r, column=j, value=h)
    c.font, c.fill, c.alignment = HEAD, HFILL, CENTER
r += 1
wc, sc = rng(L["بخش"]), rng(k["status"])
for w in WARDS:
    sm.cell(row=r, column=1, value=w).font = BODY
    sm.cell(row=r, column=2, value=f'=COUNTIF({wc},"{w}")')
    sm.cell(row=r, column=3, value=f'=COUNTIFS({wc},"{w}",{sc},"مطابق")+COUNTIFS({wc},"{w}",{sc},"مطابق، دوز مازاد")')
    sm.cell(row=r, column=4, value=f"=IFERROR(C{r}/B{r},0)").number_format = "0.0%"
    sm.cell(row=r, column=5, value=f'=SUMIF({wc},"{w}",{rng(k["value"])})').number_format = "#,##0"
    sm.cell(row=r, column=6, value=f'=SUMIF({wc},"{w}",{rng(k["unacc"])})').number_format = "#,##0"
    for j in range(1, 7):
        sm.cell(row=r, column=j).border = BOX
        if j > 1:
            sm.cell(row=r, column=j).font = BODY
    r += 1
sm.column_dimensions["A"].width = 46
for col in "BCDEF":
    sm.column_dimensions[col].width = 18

# ------------------------------------------------------------------ راهنما
g = wb.create_sheet("راهنما", 0)
g.sheet_view.rightToLeft = True
g.column_dimensions["A"].width = 110
lines = [
    ("راهنمای ممیزی پرونده‌های آلبومین — بیمارستان جم", TITLE),
    ("", BODY),
    ("هدف", BOLD),
    ("تعیین خط پایه «درصد مصرف آلبومین مطابق اندیکاسیون»، کامل بودن مستندات، ویال‌های بدون سرنوشت و ریسک "
     "کسورات بیمه، پیش از اجرای فرم جدید درخواست آلبومین؛ و تکرار همین ممیزی ۳ ماه بعد برای سنجش اثر فرم.", BODY),
    ("", BODY),
    ("نمونه‌گیری", BOLD),
    ("۱. از HIS فهرست همه بستری‌هایی را بگیرید که در دوره ممیزی آلبومین دریافت کرده‌اند (پیشنهاد: شهریور و مهر ۱۴۰۵).", BODY),
    ("۲. اولویت: بخش‌های 3A، 3B، 4A و 4B (بیشترین مصرف در بخش عمومی)، سپس ICU و I.C.U OH.", BODY),
    ("۳. از هر بخش ۳۰ پرونده به‌صورت تصادفی انتخاب کنید (اگر کمتر از ۳۰ بود، همه). با ۳۰ پرونده دقت برآورد "
     "حدود ±۱۸ درصد است و با ۱۲۰ پرونده در کل حدود ±۹ درصد.", BODY),
    ("۴. انتخاب تصادفی: فهرست را بر اساس شماره پرونده مرتب کنید و با RAND() یا هر nامین پرونده انتخاب کنید. "
     "پرونده‌ها را بر اساس «جالب بودن» انتخاب نکنید.", BODY),
    ("", BODY),
    ("نحوه تکمیل برگه «ممیزی»", BOLD),
    ("• فقط سلول‌های زرد را پر کنید. سلول‌های آبی فرمول‌اند و خودکار حساب می‌شوند.", BODY),
    ("• کد اندیکاسیون را از روی آنچه در پرونده آمده تعیین کنید، نه از روی حدس. اگر از پرونده معلوم نیست: U0.", BODY),
    ("• «معیار الزامی مستند است» فقط وقتی «بله» است که یافته یا آزمایش ستون D جدول اندیکاسیون‌ها پیش از "
     "تجویز در پرونده ثبت شده باشد.", BODY),
    ("• «نامرتبط» برای مواردی است که آن بند برای این بیمار لازم نیست (مثلاً تأیید خارج از معیار برای مورد مطابق جدول).", BODY),
    ("• ویال تحویلی را از HIS، ویال تزریق‌شده را از گزارش پرستاری/کاردکس و ویال برگشتی را از سامانه داروخانه بخوانید. "
     "اگر صفر است، صفر بنویسید، نه خالی.", BODY),
    ("• قیمت خرید هر ویال را بر اساس کد کالای تحویلی وارد کنید (مثلاً حدود ۲۱ میلیون ریال برای کد ۱۰۵۰۰۸ و "
     "۷۸٫۶ میلیون ریال برای کد ۱۲۸۶۰۹، طبق گزارش HIS ۱۴۰۵).", BODY),
    ("• سقف مرجع از جدول اندیکاسیون‌ها می‌آید. برای پلاسمافرز یا وقتی سقف برای این بیمار فرق دارد، ستون «سقف برای "
     "این بیمار» را پر کنید.", BODY),
    ("• شماره پرونده را فقط در فایل داخلی نگه دارید. در گزارش‌هایی که بیرون از واحد می‌رود، فقط ردیف را ذکر کنید.", BODY),
    ("", BODY),
    ("وضعیت‌ها", BOLD),
    ("مطابق: اندیکاسیون جدول + معیار مستند + در سقف دوز | مطابق، دوز مازاد: اندیکاسیون درست اما بیش از سقف | "
     "معیار مستند نشده: اندیکاسیون جدول ذکر شده اما شاهدش در پرونده نیست | بدون اندیکاسیون: یکی از موارد N | "
     "خارج از جدول بدون تأیید: کد X1 بدون تأیید کمیته | اندیکاسیون نامشخص: U0", BODY),
    ("", BODY),
    ("نمونه یک ردیف تکمیل‌شده (فقط برای راهنمایی؛ در برگه ممیزی وارد نکنید)", BOLD),
    ("پرونده 405123 | بخش 3B | تامین اجتماعی | 1405/06/12 | 72 kg | کد A1 | فرم: بله | دستور: بله | معیار: بله | "
     "آزمایش: بله | دوز و نقطه پایان: خیر | داروساز: بله | خارج معیار: نامرتبط | ثبت پرستاری: بله | تحویلی 6 | "
     "تزریق 6 | برگشتی 0 | قیمت 21,000,000 | سهم بیمه 0 | سهم بیمار 138,000,000 → وضعیت «مطابق»، امتیاز "
     "مستندات ۸۶٪، پیام «بررسی امکان دریافت از بیمه».", BODY),
    ("", BODY),
    ("خروجی", BOLD),
    ("برگه «خلاصه» نرخ مصرف مطابق، کامل بودن هر بند مستندات، ارزش ریالی مصرف نامطابق، ویال بدون سرنوشت و "
     "پیام‌های بیمه را به تفکیک بخش نشان می‌دهد. این اعداد مستقیماً وارد گزارش یک‌صفحه‌ای مدیریت می‌شوند.", BODY),
    ("", BODY),
    ("یادآوری: قضاوت بالینی موارد مرزی با داروساز بالینی یا پزشک معالج است. ممیزی برای اصلاح سیستم است، نه "
     "سرزنش فرد.", Font(name=F, size=10, italic=True)),
]
for i, (t, font) in enumerate(lines, 1):
    c = g.cell(row=i, column=1, value=t)
    c.font, c.alignment = font, WRAP
leg = len(lines) + 2
g.cell(row=leg, column=1, value="راهنمای رنگ: زرد = ورود داده توسط ممیز | آبی = محاسبه خودکار").font = BOLD
g.cell(row=leg, column=1).fill = INPUT

wb.save(OUT)
print(OUT)

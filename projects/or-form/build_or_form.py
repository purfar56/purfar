"""ساخت فرم ثبت روزانه دستکش معاینه و اقلام قابل دریافت اتاق عمل."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent / "or_consumption_form.xlsx"
FIRST, N = 4, 300
LAST = FIRST + N - 1
F = "Arial"
TITLE = Font(name=F, size=14, bold=True)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E7EEF7")
HFILL = PatternFill("solid", fgColor="1F4E78")
GROUP = PatternFill("solid", fgColor="D9D9D9")
RED = PatternFill("solid", fgColor="F8CBAD")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="right", readingOrder=2)
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center", readingOrder=2)

wb = Workbook()

# ---------------------------------------------------------------- پارامترها
pr = wb.active
pr.title = "پارامترها"
pr.sheet_view.rightToLeft = True
pr["A1"] = "پارامترهای فرم"
pr["A1"].font = TITLE
params = [
    ("سقف دستکش معاینه قابل ثبت برای هر عمل (جفت)", 5,
     "جزوه هتلینگ ۱۴۰۵ دانشگاه علوم پزشکی ایران؛ «برای هر عمل» به تأیید ناظر دارویی"),
    ("قیمت فروش هر عدد دستکش نیتریل (ریال)", 241920,
     "آخرین قیمت فروش در گزارش HIS شهریور ۱۴۰۵؛ با تعرفه مصوب جایگزین شود"),
    ("قیمت فروش هر عدد دستکش وینیل (ریال)", 170800,
     "آخرین قیمت فروش در گزارش HIS شهریور ۱۴۰۵؛ با تعرفه مصوب جایگزین شود"),
]
for i, h in enumerate(["پارامتر", "مقدار", "منبع"], 1):
    c = pr.cell(row=2, column=i, value=h)
    c.font, c.fill, c.alignment = HEAD, HFILL, CENTER
for r, (k, v, src) in enumerate(params, 3):
    pr.cell(row=r, column=1, value=k).font = BODY
    c = pr.cell(row=r, column=2, value=v)
    c.font, c.fill, c.number_format = Font(name=F, size=10, color="0000FF"), INPUT, "#,##0"
    pr.cell(row=r, column=3, value=src).font = Font(name=F, size=9, italic=True)
pr.column_dimensions["A"].width = 46
pr.column_dimensions["B"].width = 14
pr.column_dimensions["C"].width = 70
CAP = "'پارامترها'!$B$3"
P_NIT = "'پارامترها'!$B$4"
P_VIN = "'پارامترها'!$B$5"

# ---------------------------------------------------------------- ثبت روزانه
ws = wb.create_sheet("ثبت روزانه")
ws.sheet_view.rightToLeft = True
ws["A1"] = "فرم ثبت مصرف دستکش معاینه و اقلام قابل دریافت — اتاق عمل بیمارستان جم"
ws["A1"].font = TITLE
ws["A2"] = ("هر ردیف یک عمل. دستکش‌های داخل پک جنرال یا پک اعمال جراحی را ننویسید؛ فقط دستکش معاینه‌ای که "
            "جدا از پک باز شده است. سلول‌های زرد را پر کنید.")
ws["A2"].font = Font(name=F, size=9, italic=True)

YN = "yn"
COLS = [
    ("ردیف", "calc", 5, "row"),
    ("تاریخ عمل", "in", 11, None),
    ("شماره پرونده", "in", 12, None),
    ("شماره اتاق", "in", 7, None),
    ("نام عمل", "in", 22, None),
    ("جراح", "in", 14, None),
    ("پرستار سیار", "in", 14, None),
    ("دستکش نیتریل (جفت)", "in", 9, "num"),
    ("دستکش وینیل (جفت)", "in", 9, "num"),
    ("جمع دستکش (جفت)", "calc", 8, "total"),
    ("قابل ثبت (جفت)", "calc", 8, "billable"),
    ("مازاد بر سقف (جفت)", "calc", 8, "excess"),
    ("چسب خارج از پک (عدد)", "in", 9, "num"),
    ("سرنگ خارج از پک (عدد)", "in", 9, "num"),
    ("گاز/باند خارج از پک (عدد)", "in", 9, "num"),
    ("ثبت در HIS", "in", 8, YN),
    ("ثبت‌کننده HIS", "in", 12, None),
    ("ارزش دستکش قابل ثبت (ریال)", "calc", 14, "value"),
    ("وضعیت", "calc", 18, "status"),
    ("توضیح", "in", 22, None),
]
L = {name: get_column_letter(i) for i, (name, *_) in enumerate(COLS, 1)}
f_file = L["شماره پرونده"]
nit, vin = L["دستکش نیتریل (جفت)"], L["دستکش وینیل (جفت)"]
tot, bil, exc = L["جمع دستکش (جفت)"], L["قابل ثبت (جفت)"], L["مازاد بر سقف (جفت)"]
his = L["ثبت در HIS"]


def formula(key, r):
    blank = f'{f_file}{r}=""'
    if key == "row":
        return f'=IF({blank},"",ROW()-{FIRST - 1})'
    if key == "total":
        return f'=IF({blank},"",N({nit}{r})+N({vin}{r}))'
    if key == "billable":
        return f'=IF({blank},"",MIN({tot}{r},{CAP}))'
    if key == "excess":
        return f'=IF({blank},"",MAX(0,{tot}{r}-{CAP}))'
    if key == "value":
        # قابل ثبت به نسبت نیتریل و وینیل تقسیم می‌شود؛ هر جفت = ۲ عدد
        return (f'=IF(OR({blank},N({tot}{r})=0),"",{bil}{r}*2*'
                f'(N({nit}{r})*{P_NIT}+N({vin}{r})*{P_VIN})/{tot}{r})')
    if key == "status":
        return (f'=IF({blank},"",IF({his}{r}<>"بله","ثبت نشده در HIS",'
                f'IF(N({exc}{r})>0,"ثبت‌شده؛ مازاد بر سقف","ثبت‌شده")))')
    raise KeyError(key)


for j, (name, kind, width, spec) in enumerate(COLS, 1):
    col = get_column_letter(j)
    c = ws.cell(row=3, column=j, value=name)
    c.font, c.fill, c.alignment, c.border = HEAD, HFILL, CENTER, BOX
    ws.column_dimensions[col].width = width
    for r in range(FIRST, LAST + 1):
        cell = ws.cell(row=r, column=j)
        cell.font, cell.border, cell.alignment = BODY, BOX, CENTER
        if kind == "calc":
            cell.value = formula(spec, r)
            cell.fill = CALC
        else:
            cell.fill = INPUT
    if spec == "value":
        for r in range(FIRST, LAST + 1):
            ws[f"{col}{r}"].number_format = "#,##0"
ws.row_dimensions[3].height = 60
ws.freeze_panes = f"{L['نام عمل']}{FIRST}"
ws[f"{exc}3"].comment = Comment("مازاد بر ۵ جفت در HIS ثبت نشود؛ بیمه آن را کسر می‌کند.", "ناظر دارویی")
ws[f"{bil}3"].comment = Comment("همین عدد (حداکثر ۵ جفت) در HIS ثبت شود.", "ناظر دارویی")

lists = wb.create_sheet("فهرست")
lists["A1"], lists["A2"] = "بله", "خیر"
lists.sheet_state = "hidden"
dv = DataValidation(type="list", formula1="'فهرست'!$A$1:$A$2", allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{his}{FIRST}:{his}{LAST}")
num = DataValidation(type="whole", operator="between", formula1="0", formula2="50", allow_blank=True)
num.error = "عدد صحیح بین ۰ و ۵۰"
ws.add_data_validation(num)
for name, _, _, spec in COLS:
    if spec == "num":
        num.add(f"{L[name]}{FIRST}:{L[name]}{LAST}")
ws.conditional_formatting.add(f"{exc}{FIRST}:{exc}{LAST}",
                              FormulaRule(formula=[f'AND({exc}{FIRST}<>"",N({exc}{FIRST})>0)'], fill=RED))
st = L["وضعیت"]
ws.conditional_formatting.add(f"{st}{FIRST}:{st}{LAST}",
                              FormulaRule(formula=[f'{st}{FIRST}="ثبت نشده در HIS"'], fill=RED))

ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "3:3"
ws.print_options.gridLines = True

# ---------------------------------------------------------------- خلاصه
sm = wb.create_sheet("خلاصه")
sm.sheet_view.rightToLeft = True
sm["A1"] = "خلاصه ثبت دستکش اتاق عمل"
sm["A1"].font = TITLE
S = "'ثبت روزانه'!"


def rng(col):
    return f"{S}${col}${FIRST}:${col}${LAST}"


rows = [
    ("تعداد عمل ثبت‌شده در فرم", f"=COUNTA({rng(f_file)})", "0"),
    ("عمل‌هایی که دستکش در HIS ثبت شده", f'=COUNTIF({rng(his)},"بله")', "0"),
    ("نرخ ثبت در HIS", "=IFERROR(B4/B3,0)", "0.0%"),
    ("جمع دستکش معاینه مصرفی (جفت)", f"=SUM({rng(tot)})", "#,##0"),
    ("دستکش قابل ثبت (جفت)", f"=SUM({rng(bil)})", "#,##0"),
    ("مازاد بر سقف — غیرقابل ثبت (جفت)", f"=SUM({rng(exc)})", "#,##0"),
    ("عمل‌هایی با مصرف بیش از ۵ جفت", f'=COUNTIF({rng(exc)},">0")', "0"),
    ("میانگین دستکش به ازای هر عمل (جفت)", "=IFERROR(B6/B3,0)", "0.0"),
    ("ارزش دستکش قابل ثبت (ریال)", f"=SUM({rng(L['ارزش دستکش قابل ثبت (ریال)'])})", "#,##0"),
    ("ارزش قابل ثبتی که در HIS ثبت نشده (ریال)",
     f'=SUMIF({rng(his)},"<>بله",{rng(L["ارزش دستکش قابل ثبت (ریال)"])})', "#,##0"),
    ("چسب خارج از پک (عدد)", f"=SUM({rng(L['چسب خارج از پک (عدد)'])})", "#,##0"),
    ("سرنگ خارج از پک (عدد)", f"=SUM({rng(L['سرنگ خارج از پک (عدد)'])})", "#,##0"),
    ("گاز/باند خارج از پک (عدد)", f"=SUM({rng(L['گاز/باند خارج از پک (عدد)'])})", "#,##0"),
]
for i, (k, f, fmt) in enumerate(rows, 3):
    a = sm.cell(row=i, column=1, value=k)
    b = sm.cell(row=i, column=2, value=f)
    a.font, b.font, a.alignment = BODY, BOLD, WRAP
    a.border = b.border = BOX
    b.number_format = fmt
sm.column_dimensions["A"].width = 46
sm.column_dimensions["B"].width = 20
note = sm.cell(row=len(rows) + 4, column=1,
               value="ارزش‌ها با قیمت‌های برگه «پارامترها» حساب می‌شوند. چسب، سرنگ و گاز/باند خارج از پک هم "
                     "در اتاق عمل قابل دریافت‌اند (طبق جزوه هتلینگ) و فعلاً فقط شمرده می‌شوند.")
note.font, note.alignment = Font(name=F, size=9, italic=True), WRAP

# ---------------------------------------------------------------- راهنما
g = wb.create_sheet("راهنما", 0)
g.sheet_view.rightToLeft = True
g.column_dimensions["A"].width = 110
lines = [
    ("راهنمای فرم ثبت دستکش و اقلام قابل دریافت اتاق عمل", TITLE),
    ("", BODY),
    ("چرا این فرم؟", BOLD),
    ("دستکش معاینه در اتاق عمل تا ۵ جفت برای هر عمل قابل دریافت از بیمه است. در فروردین تا شهریور ۱۴۰۵، "
     "۵۹٬۰۰۰ عدد دستکش معاینه به اتاق عمل رفته و هیچ‌کدام روی پرونده ثبت نشده است.", BODY),
    ("", BODY),
    ("چه کسی، چه زمانی", BOLD),
    ("• پرستار سیار در پایان هر عمل تعداد جفت دستکش معاینه نیتریل و وینیل را که جدا از پک باز شده می‌نویسد.", BODY),
    ("• منشی یا پرستار سیار «قابل ثبت (جفت)» را (حداکثر ۵ جفت) همان روز در HIS روی پرونده ثبت می‌کند و "
     "ستون «ثبت در HIS» را «بله» می‌زند.", BODY),
    ("• سرپرستار اتاق عمل هر روز ردیف‌های قرمز («ثبت نشده در HIS») را پیگیری می‌کند.", BODY),
    ("• ناظر دارویی هر هفته برگه «خلاصه» را بررسی و نتیجه را در گزارش دوشنبه می‌آورد.", BODY),
    ("", BODY),
    ("قواعد", BOLD),
    ("• دستکش داخل پک جنرال یا پک اعمال جراحی نوشته نشود؛ پک جزء فنی اتاق عمل است و جداگانه قابل دریافت نیست.", BODY),
    ("• دستکش جراحی استریل جزو این فرم نیست و طبق روال فعلی ثبت می‌شود.", BODY),
    ("• بیش از ۵ جفت در HIS ثبت نشود. مصرف واقعی بیشتر از ۵ جفت را در فرم بنویسید تا مصرف واقعی معلوم باشد، "
     "ولی در HIS فقط ۵ جفت ثبت کنید.", BODY),
    ("• چسب، سرنگ و گاز/باند هم اگر خارج از پک مصرف شود در اتاق عمل قابل دریافت است؛ تعدادش را بنویسید.", BODY),
    ("", BODY),
    ("نمونه یک ردیف (در برگه ثبت وارد نکنید)", BOLD),
    ("1405/08/03 | پرونده 512340 | اتاق 2 | کوله‌سیستکتومی لاپاراسکوپیک | نیتریل 4 جفت | وینیل 2 جفت → "
     "جمع 6، قابل ثبت 5، مازاد 1 | چسب 2 | ثبت در HIS: بله → وضعیت «ثبت‌شده؛ مازاد بر سقف»", BODY),
    ("", BODY),
    ("راهنمای رنگ: زرد = ورود داده | آبی = محاسبه خودکار | قرمز = نیاز به پیگیری", BOLD),
]
for i, (t, font) in enumerate(lines, 1):
    c = g.cell(row=i, column=1, value=t)
    c.font, c.alignment = font, WRAP

wb.save(OUT)
print(OUT)

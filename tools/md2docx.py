import re, sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
src, dst = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles['Normal']; st.font.name = 'B Nazanin'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:cs'), 'B Nazanin')
def rtl(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:bidi'); pPr.append(b)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
def add_runs(p, text):
    for part in re.split(r'(\*\*[^*]+\*\*)', text):
        if not part: continue
        bold = part.startswith('**')
        r = p.add_run(part.strip('*') if bold else part); r.bold = bold
        r.font.rtl = True; r._r.get_or_add_rPr()
        rf = r._r.rPr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); r._r.rPr.append(rf)
        rf.set(qn('w:cs'), 'B Nazanin')
        if bold:
            bcs = OxmlElement('w:bCs'); r._r.rPr.append(bcs)
lines = open(src, encoding='utf-8').read().split('\n')
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if ln.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'-+', c) for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        tblPr = t._tbl.tblPr; b = OxmlElement('w:bidiVisual'); tblPr.append(b)
        for ri, row in enumerate(rows):
            for ci, c in enumerate(row):
                cell = t.cell(ri, ci); p = cell.paragraphs[0]; rtl(p)
                add_runs(p, f'**{c}**' if ri == 0 and not c.startswith('**') else c)
                for r in p.runs: r.font.size = Pt(10)
        continue
    if ln == '---': i += 1; continue
    if ln.strip() == '': i += 1; continue
    m = re.match(r'^(\d+)\.\s+(.*)', ln); bullet = ln.startswith('- ')
    text = m.group(2) if m else (ln[2:] if bullet else ln)
    text = text.rstrip('\\')
    p = doc.add_paragraph(); rtl(p)
    if m: text = f'{m.group(1)}. ' + text
    if bullet: text = '• ' + text
    add_runs(p, text)
    i += 1
doc.save(dst)
